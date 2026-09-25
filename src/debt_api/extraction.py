"""Extraction of calculation parameters from Brazilian court documents with a language model.

Non-negotiable principle: **the model never does the arithmetic**. It reads the judgment and
returns structured parameters that the lawyer reviews before the engine calculates. Every
field carries its source excerpt (a literal quote) and a confidence; what the document does
not state comes back as null ("not identified"), never invented.

Two modes:
* **Live**: requires ANTHROPIC_API_KEY. The model reads the PDF(s) natively (scanned ones
  included) and returns JSON validated by Pydantic through structured outputs
  (`client.messages.parse`).
* **Simulated**: without a key, returns the extraction of reference case 01 (flagged
  `simulated=true`), with fictional parties, for development and for demonstrating the UI
  without API cost.

Model: env EXTRACTION_MODEL (default claude-sonnet-5; move up to a larger model if the
measured accuracy asks for it).
"""

from __future__ import annotations

import base64
import io
import os

from pydantic import BaseModel, Field

from debt_api.demo import DEMO_CASE_NUMBER, DEMO_CREDITOR, DEMO_DEBTOR

DEFAULT_MODEL = "claude-sonnet-5"
MAX_BYTES_PER_FILE = 20 * 1024 * 1024  # API limit: 32 MB per request
# Margin under the serverless ceiling (60 s on the hobby tier; 50 here). On a paid tier, raise
# maxDuration in vercel.json (up to 300) and set EXTRACTION_TIMEOUT (e.g. 280): large
# multi-document cases (3 PDFs measured at 68 s) do not fit the hobby tier.
TIMEOUT_SECONDS = float(os.environ.get("EXTRACTION_TIMEOUT", "50"))


def extraction_mode() -> str:
    return "live" if os.environ.get("ANTHROPIC_API_KEY") else "simulated"


class ExtractedField(BaseModel):
    """One extracted parameter with provenance for the human review."""

    value: str | None = Field(
        description="Extracted value, or null when the document does not state it ('not identified')"
    )
    quote: str | None = Field(
        description="Short LITERAL excerpt of the document the value came from (max ~200 chars)"
    )
    document: str | None = Field(
        default=None,
        description="File name (exactly as in the attachment title) the quote came from, e.g. 'judgment.pdf'. "
        "null when the value did not come from a document.",
    )
    confidence: float = Field(description="Confidence from 0 to 1 in the extraction of this field")


class ExtractedInstallment(BaseModel):
    amount: str = Field(description="Monetary amount as a decimal with a dot, e.g. '134000.00'")
    on: str = Field(description="Date of the amount in ISO (YYYY-MM-DD)")
    description: str = Field(
        description="Short description, e.g. 'Economic benefit', 'Installment 3'"
    )
    quote: str | None = Field(description="Literal excerpt it came from")
    document: str | None = Field(
        default=None, description="File name the installment came from (e.g. 'history.xlsx')"
    )
    confidence: float


class ExtractedDeduction(BaseModel):
    """Deduction by AMOUNT (an abatement without its own adjustment date), e.g. an assessed
    usage fee, outstanding condominium charges, brokerage withheld when the decision orders it."""

    description: str = Field(
        description="Short description, e.g. 'Outstanding condominium charges'"
    )
    amount: str = Field(description="Amount as a decimal with a dot, e.g. '1097.66'")
    quote: str | None = Field(description="Literal excerpt it came from")
    document: str | None = Field(default=None, description="File name the deduction came from")
    confidence: float


class ExtractedEntry(BaseModel):
    """Court fee / expense / discount WITH A DATE: the engine adjusts it from its own date (no
    interest). The amount is always the ORIGINAL one on the date of the payment, NEVER the
    already-adjusted amount a calculation spreadsheet may show next to it."""

    kind: str = Field(
        description="'court_fee' (court fees, ADD), 'expense' (procedural expenses, ADD) or 'discount' "
        "(discount/abatement/recognised partial payment, SUBTRACTS)"
    )
    description: str = Field(description="Short description, e.g. 'Court fee', 'Abatement'")
    amount: str = Field(
        description="ORIGINAL amount on the date of the entry, decimal with a dot, e.g. '10003.47'. "
        "NEVER the updated/adjusted amount"
    )
    on: str = Field(description="Date of the payment in ISO (YYYY-MM-DD)")
    quote: str | None = Field(description="Literal excerpt it came from")
    document: str | None = Field(default=None, description="File name the entry came from")
    confidence: float


class RecognisedDocument(BaseModel):
    """Classification of each attached document: transparency for the human review."""

    file_name: str = Field(description="Name of the attached file this refers to")
    kind: str = Field(
        description="Kind of document: 'judgment' (sentença), 'appellate decision' (acórdão), 'single-judge decision', "
        "'motion for clarification' (embargos de declaração), 'interlocutory decision', 'order' (despacho), "
        "'settlement approval' (homologação de acordo), 'petition', 'contract', 'payment history', "
        "'calculation spreadsheet', 'certificate' or 'other'"
    )
    on: str | None = Field(
        description="Date of the document in ISO (YYYY-MM-DD) when identifiable; otherwise null"
    )
    role: str = Field(
        description="One line: what this document contributes to the calculation (e.g. 'sets the index and the "
        "interest', 'reverses the judgment as to the fees', 'source of the paid installments')"
    )


class ExtractionResult(BaseModel):
    """Mirrors the fields of the form / CalculationIn: it is what pre-fills the screen."""

    documents: list[RecognisedDocument] = Field(
        description="One entry PER attached file, in order, classifying the document"
    )
    prevailing_decision: ExtractedField = Field(
        description="Which document GOVERNS the calculation (the most recent and hierarchically applicable "
        "decision) and why, e.g. 'appellate decision of 12/05/2025, which partially reversed the judgment as "
        "to the interest'. Quote = the excerpt of the operative part that proves it"
    )
    settlement: ExtractedField = Field(
        description="Terms of the approved settlement (amount, installments, charges), ONLY when a decision "
        "EXPLICITLY approves a settlement/transaction; in that case the settlement becomes the base of the "
        "calculation and prevails over the merits. null when there is no express approval"
    )
    case_number: ExtractedField
    creditor: ExtractedField
    debtor: ExtractedField
    installments: list[ExtractedInstallment] = Field(
        description="Amounts to update (award, paid installments, awarded fees...)"
    )
    index: ExtractedField = Field(
        description="Adjustment index set by the decision: 'tjsp' (TJSP practical table: INPC until 07/2024 "
        "and IPCA-15 from 08/2024), 'tjdft' (INPC until 08/2024 and full IPCA afterwards, the standard after "
        "Law 14,905 outside São Paulo), 'inpc', 'ipca', 'ipca15', 'igpm' or 'igpdi'"
    )
    interest_regime: ExtractedField = Field(
        description="'legal' (Civil Code art. 406: 1% per month -> Legal Rate on 30/08/2024), 'legal_rate', "
        "'fixed' or 'none'"
    )
    interest_fixed_rate: ExtractedField = Field(
        description="% per month when the decision sets its own rate (regime 'fixed'), e.g. '1.0'"
    )
    interest_start: ExtractedField = Field(
        description="Start of the interest in ISO (service of process, final judgment, harmful event). null "
        "when it depends on data that is not in the documents"
    )
    penalty_pct: ExtractedField = Field(
        description="% of the contractual penalty / penalty clause / rescission penalty DEDUCTED from the refund "
        "(e.g. '25' when the decision withholds 25% of the paid amounts). NOT a fine that ADDS to the debt: "
        "that one goes in 'late_fine_pct'"
    )
    late_fine_pct: ExtractedField = Field(
        description="% of the late/contractual fine that ADDS to the debt, computed over the adjusted amount "
        "(e.g. 'a 10% addition for the fine' in a spreadsheet, a condominium fine under Civil Code art. 1,336 "
        "§1, a contractual late fine). Do NOT confuse with the contractual penalty (which deducts) nor with "
        "the art. 523 fine (enforcement of a judgment)"
    )
    title_fee_pct: ExtractedField = Field(
        description="% of attorney fees provided for in the TITLE itself being enforced (a contract / "
        "debt-acknowledgement clause with the debtor's consent, e.g. 'attorney fees of 20% over the total "
        "amount of the debt, to which the debtor consents'): they ADD to the debt. Different from "
        "'awarded_fee_pct' (set in a COURT decision). If the clause applies 'over the total amount of the "
        "debt' (with the fine), record that in 'notes'"
    )
    deductions: list[ExtractedDeduction] = Field(
        description="Deductions with an ALREADY ASSESSED AMOUNT and WITHOUT their own adjustment date "
        "(quantified usage fee, condominium charges, brokerage to abate...). Empty when there are none"
    )
    entries: list[ExtractedEntry] = Field(
        description="Court fees, procedural expenses and discounts/abatements WITH A DATE: one item per entry, "
        "with the ORIGINAL amount and the date. Empty when there are none"
    )
    deductions_mentioned: str = Field(
        description="ONLY deductions mentioned WITHOUT an assessable amount in the documents (e.g. 'usage fee "
        "of 0.5% per month to be assessed in liquidation'), with the criterion of each one; short free text; "
        "empty when there are none. What has an amount and a date goes in 'entries'; what has an amount "
        "without a date goes in 'deductions'"
    )
    awarded_fee_pct: ExtractedField = Field(
        description="% of the fees awarded by the court (CPC art. 85) set in the decision; with reciprocal "
        "pro-rata awards, the effective % of each party, explaining in notes"
    )
    awarded_fee_base: ExtractedField = Field(
        description="Base of the awarded fees, value MANDATORILY 'net' or 'claim_value' (no other text): 'net' "
        "when they accrue on the award / economic benefit assessed by the calculation (the case of '% over "
        "the economic benefit': the benefit is the entered installments); 'claim_value' ONLY when the "
        "decision EXPRESSLY orders the calculation over the claim value (then also fill 'claim_value')"
    )
    claim_value: ExtractedField = Field(
        description="Claim value as a decimal with a dot (e.g. '48990.00'): fill ONLY when "
        "awarded_fee_base='claim_value' and the amount appears in the documents; otherwise null"
    )
    awarded_fees_are_the_claim: ExtractedField = Field(
        description="'yes' when the enforcement collects ONLY the awarded fees (the enforced credit is the "
        "fees themselves; the amounts/benefit are just the base of the calculation, e.g. a petition titled "
        "'enforcement of judgment: awarded attorney fees' asking for the payment of the fee amount). null "
        "when the enforcement is of the main award (with or without fees added)"
    )
    fine_523_pct: ExtractedField = Field(
        description="% of the fine of CPC art. 523 §1, when requested"
    )
    fee_523_pct: ExtractedField = Field(
        description="% of the fees of CPC art. 523 §1, when requested"
    )
    notes: str = Field(
        description="Alerts for the lawyer: ambiguities, conflicting decisions, fields that require a document "
        "that was not attached (e.g. the date of service is in the case history)"
    )


INSTRUCTIONS = """You extract court-calculation parameters from Brazilian court documents (written in \
Portuguese) to feed a deterministic liquidation engine. You NEVER compute amounts: you only extract what \
is written. Field values follow the enumerations of the output schema; free text (quotes, descriptions, \
notes) may stay in the document's language.

SUPREME RULE, ABSOLUTE FIDELITY TO THE SOURCE (above everything):
You may only state a fact if it is LITERALLY WRITTEN in the provided documents. It is STRICTLY FORBIDDEN \
to invent, infer, deduce or "complete" amounts, percentages, dates, indices, decisions or settlements. If \
something is not explicit, IT DOES NOT EXIST for you: return null ("not identified") and explain in \
'notes'. A single invented datum makes the whole calculation challengeable. Do NOT carry facts from one \
case into another.

DOCUMENTS AND HIERARCHY (SEVERAL documents of different kinds may come):
Classify EACH attached file in 'documents': judgment (sentença), appellate decision (acórdão), single-judge \
decision, motion for clarification (embargos de declaração), interlocutory decision, order (despacho), \
settlement approval (homologação de acordo), petition, contract, payment history, calculation spreadsheet, \
certificate or other. One file may contain more than one document (concatenated case files): classify by \
the dominant content and explain in 'role'. Then apply the hierarchy to find the decision that GOVERNS the \
calculation ('prevailing_decision'):
1. The operative part (dispositivo) prevails over the reasoning, in any document.
2. The MOST RECENT decision of a higher instance prevails ON THE POINT IT CHANGED: an appellate decision \
replaces the judgment on the reversed points (on the others, the judgment stands); a single-judge decision \
or a decision of a superior court (STJ/STF) prevails over the appellate decision; a motion for \
clarification only changes the result when it has a MODIFYING effect that was granted.
3. SETTLEMENT/TRANSACTION: only consider a settlement when a decision EXPLICITLY approves it. NEVER infer a \
settlement from deposit receipts, petitions not yet approved or generic mentions. With an express \
approval, the settlement PREVAILS over the merits and becomes the base of the calculation: extract in \
'settlement' the terms that APPEAR (amount, installments, due dates, charges, fees) and align \
'installments' and the other fields with the terms of the settlement.
4. Always extract the FINAL regime in force after applying 1 to 3. If the documents contain no decision \
able to ground the calculation (only petitions/orders, an appeal pending judgment), say so in \
'prevailing_decision' (value null) and in 'notes'.

CROSS USE OF THE DOCUMENTS: each kind of document feeds different fields: the prevailing decision sets the \
index, the interest, the penalty and the fees; the history/spreadsheet provides the paid installments and \
their dates (one entry per installment); the history usually carries the date of service; the contract \
grounds the penalty clause and the rates. Combine them, always with the quote of the document each datum \
came from.

SPREADSHEETS (Excel/CSV): attachments marked "[ATTACHED SPREADSHEET: ...]" are spreadsheets converted to \
tabular text (cells separated by " | "). Typically they are a payment history or a calculation statement: \
classify them in 'documents' as 'payment history' or 'calculation spreadsheet' and extract each payment \
line as an installment (amount + date). Amounts may come in the Brazilian format ("1.850,00"): normalise to \
a decimal with a dot. The quote of data coming from a spreadsheet is the tabular line itself.

EXTRACTION RULES:
- Absent = null. NEVER invent, estimate or complete an amount, date or percentage that is not written in \
the documents. A field without an explicit textual basis -> value null and low confidence.
- Every extraction carries 'quote': the short LITERAL excerpt of the document that supports it, and \
'document': the NAME of the file (equal to the attachment title) the excerpt came from.
- Dates in ISO (YYYY-MM-DD). Monetary amounts as a decimal with a dot ('134000.00').
- Index: map the decision's determination to 'tjsp' (the TJSP practical table, INPC -> IPCA-15; use it \
when the decision orders the TJSP table or is from a São Paulo court ordering adjustment by 'official \
indices'), 'tjdft' (INPC -> full IPCA of Law 14,905/2024; use it when the decision orders INPC and Law \
14,905 applies to the period, or when it orders 'official indices' outside São Paulo), \
'inpc'/'ipca'/'ipca15'/'igpm'/'igpdi' (an expressly named single index). In doubt, pick the most likely \
one and lower the confidence, explaining in 'notes'.
- Interest: 'legal' when the decision orders legal interest / 1% per month of Civil Code art. 406 (the \
engine applies the 1% -> Legal Rate timeline automatically); 'fixed' only for a contract's own rate; \
'none' when expressly excluded.
- Start of the interest: service of process, final judgment or harmful event as the decision says. If the \
decision says 'from the service of process' but the date of service is not in the attached documents, \
value null and explain in 'notes' that the history with the date is missing.
- Multiple amounts/installments: one entry per amount with its own date.
- 'installments' are ONLY the amounts that make up the credit to update: in a refund, the amounts \
ACTUALLY PAID (with the payment/receipt date). NEVER include as an installment: the negotiated price of \
the contract, unpaid or future installments, nor sums the decision orders NOT to refund (e.g. \
non-refundable brokerage: cite it in 'deductions_mentioned'). If the history has due-date AND receipt \
columns, use the RECEIPT date and amount.
- DATED COURT FEES, EXPENSES AND DISCOUNTS -> 'entries' (they are NOT installments): court fees and \
procedural expenses add to the debt; a discount/abatement/recognised partial payment subtracts. One item \
per entry with kind ('court_fee'/'expense'/'discount'), description, DATE and the ORIGINAL amount on the \
date. Calculation spreadsheets usually show each entry in two columns, the original AND the updated \
amount: ALWAYS extract the original (the engine adjusts from the date); extracting the adjusted one would \
double the adjustment.
- DEDUCTIONS WITH AN AMOUNT but WITHOUT their own adjustment date (an already quantified usage fee, \
outstanding condominium charges, brokerage the decision orders to abate) -> 'deductions' (description + \
amount). Deductions WITHOUT an assessable amount -> 'deductions_mentioned' (text).
- A FINE THAT ADDS to the debt (a late/contractual fine added to the total, e.g. 'a 10% addition for the \
fine' in a spreadsheet, a condominium fine) -> 'late_fine_pct'. It is NOT the contractual penalty \
('penalty_pct' DEDUCTS from the refund; using it would flip the sign) and it is only 'fine_523_pct' when \
the document EXPRESSLY ties it to CPC art. 523 §1 (enforcement of a judgment). If the document says the \
fees do not accrue on the fine ('not applicable over the fine'), record that in 'notes'.
- FEES: those set in a COURT decision go in 'awarded_fee_pct'; those provided for in the TITLE itself \
being enforced (a clause with the debtor's consent) go in 'title_fee_pct'; they may coexist. Fees of the \
contract between the creditor and their own law firm are NOT part of the debt: cite them only in 'notes'.
- Enforcement of the awarded fees ONLY (the petition collects the fee amount; the award/benefit is just the \
base of the calculation): set awarded_fees_are_the_claim='yes' and enter the BASE (benefit/award) in \
'installments'; the system will do base -> fees -> art. 523 on the fees. Do NOT confuse with the \
enforcement of the main award.
- 'notes' is the channel for uncertainty: list ambiguities and what the lawyer must check."""


def simulated_extraction(file_names: list[str] | None = None) -> dict:
    """The extraction of reference case 01 with fictional parties, for development without an
    API key and for the public demo."""
    names = file_names or ["document.pdf"]
    result = ExtractionResult(
        documents=[
            RecognisedDocument(
                file_name=name,
                kind="not analysed (simulated mode)",
                on=None,
                role="demonstration: the content was not read; set the API key to enable live extraction",
            )
            for name in names
        ],
        prevailing_decision=ExtractedField(
            value="enforcement judgment: the only decision in the documents (simulated)",
            quote="JULGO PROCEDENTE o pedido de cumprimento de sentença",
            confidence=0.9,
        ),
        settlement=ExtractedField(value=None, quote=None, confidence=0.95),
        case_number=ExtractedField(
            value=DEMO_CASE_NUMBER, quote=f"Autos: {DEMO_CASE_NUMBER}", confidence=0.99
        ),
        creditor=ExtractedField(
            value=DEMO_CREDITOR, quote=f"Exequentes: {DEMO_CREDITOR}", confidence=0.98
        ),
        debtor=ExtractedField(
            value=DEMO_DEBTOR, quote=f"Executada: {DEMO_DEBTOR}", confidence=0.98
        ),
        installments=[
            ExtractedInstallment(
                amount="134000.00",
                on="2022-03-02",
                description="Economic benefit (base of the fees)",
                quote="o proveito econômico da Executada foi de R$ 134.000,00 (...) o ajuizamento da ação ocorreu em 02/03/2022",
                confidence=0.95,
            )
        ],
        index=ExtractedField(
            value="tjdft",
            quote="considerando o INPC até 30/08/2024 e IPCA a partir de 01/09/2024",
            confidence=0.95,
        ),
        interest_regime=ExtractedField(
            value="legal",
            quote="incidência de juros moratórios conforme taxa legal, a partir do trânsito em julgado",
            confidence=0.93,
        ),
        interest_fixed_rate=ExtractedField(value=None, quote=None, confidence=0.9),
        interest_start=ExtractedField(
            value="2026-05-15",
            quote="transitou em julgado o acórdão no dia 15/05/2026",
            confidence=0.95,
        ),
        penalty_pct=ExtractedField(value=None, quote=None, confidence=0.9),
        late_fine_pct=ExtractedField(value=None, quote=None, confidence=0.9),
        title_fee_pct=ExtractedField(value=None, quote=None, confidence=0.9),
        deductions=[],
        entries=[],
        deductions_mentioned="",
        awarded_fee_pct=ExtractedField(
            value="10",
            quote="honorários foram arbitrados em 10% sobre o proveito econômico",
            confidence=0.9,
        ),
        awarded_fee_base=ExtractedField(
            value="net", quote="10% sobre o proveito econômico", confidence=0.85
        ),
        awarded_fees_are_the_claim=ExtractedField(
            value="yes",
            quote="CUMPRIMENTO DE SENTENÇA (HONORÁRIOS ADVOCATÍCIOS SUCUMBENCIAIS): "
            "pagamento voluntário no valor de R$ 16.236,33",
            confidence=0.9,
        ),
        claim_value=ExtractedField(value=None, quote=None, confidence=0.9),
        fine_523_pct=ExtractedField(
            value="10",
            quote="ser acrescida multa percentual de 10% (dez por cento) sobre o montante total do débito",
            confidence=0.9,
        ),
        fee_523_pct=ExtractedField(
            value="10",
            quote="bem como honorários advocatícios também no percentual de 10% (dez por cento), conforme Artigo 523, §1º",
            confidence=0.9,
        ),
        notes="SIMULATED MODE: a demonstration extraction based on reference case 01 with fictional parties; "
        "the content of the uploaded file was not read. Set ANTHROPIC_API_KEY to enable live extraction. "
        "The art. 523 fine and fees are conditional on the absence of voluntary payment within 15 days.",
    )
    return {"simulated": True, "model": None, "extraction": result.model_dump()}


# Spreadsheets (Excel/CSV) are not a native document type of the API: we convert them to
# tabular text on the server and send them as a text attachment. Limits keep a giant
# spreadsheet from blowing up the input tokens.
MAX_SPREADSHEET_ROWS = 400  # per sheet
MAX_SPREADSHEET_COLUMNS = 25


def _cell(value) -> str:
    if value is None:
        return ""
    if hasattr(value, "isoformat"):  # openpyxl datetime/date
        return value.isoformat()[:10]
    return str(value)


def spreadsheet_to_text(name: str, data: bytes) -> str:
    """Converts .xlsx/.csv into tabular text (one line per row, cells separated by ' | '), with
    an explicit warning when truncating."""
    parts: list[str] = [f"[ATTACHED SPREADSHEET: {name}, converted to tabular text]"]
    if name.lower().endswith(".csv"):
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = data.decode("latin-1", errors="replace")
        lines = text.splitlines()
        parts.extend(lines[:MAX_SPREADSHEET_ROWS])
        if len(lines) > MAX_SPREADSHEET_ROWS:
            parts.append(f"[... truncated: {len(lines) - MAX_SPREADSHEET_ROWS} lines omitted]")
        return "\n".join(parts)

    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    for sheet in wb.worksheets:
        parts.append(f"\n== SHEET: {sheet.title} ==")
        emitted = 0
        truncated = False
        for row in sheet.iter_rows(values_only=True):
            cells = [_cell(v) for v in row[:MAX_SPREADSHEET_COLUMNS]]
            if not any(c.strip() for c in cells):
                continue  # skips fully empty rows
            if emitted >= MAX_SPREADSHEET_ROWS:
                truncated = True
                break
            parts.append(" | ".join(cells).rstrip(" |"))
            emitted += 1
        if truncated:
            parts.append(f"[... truncated at {MAX_SPREADSHEET_ROWS} rows in this sheet]")
    wb.close()
    return "\n".join(parts)


def _is_pdf(data: bytes) -> bool:
    return data.startswith(b"%PDF")


def _normalise(extracted: ExtractionResult) -> ExtractionResult:
    """Deterministic normalisation of closed-domain fields (the model sometimes returns free
    text, e.g. base 'updated economic benefit')."""
    base = extracted.awarded_fee_base
    if base.value and base.value not in ("net", "claim_value"):
        base.value = (
            "claim_value"
            if "claim" in base.value.lower() or "causa" in base.value.lower()
            else "net"
        )
    claim = extracted.awarded_fees_are_the_claim
    if claim.value:
        claim.value = "yes" if claim.value.strip().lower() in ("yes", "sim", "true") else None
    for entry in extracted.entries:
        kind = entry.kind.strip().lower()
        if kind not in ("court_fee", "expense", "discount"):
            if "fee" in kind or "cust" in kind:
                entry.kind = "court_fee"
            elif any(t in kind for t in ("discount", "abat", "payment", "descont", "pagamento")):
                entry.kind = "discount"
            else:
                entry.kind = "expense"
    return extracted


def extract_with_model(files: list[tuple[str, bytes]]) -> dict:
    """Live extraction: native PDFs (scanned ones included) + spreadsheets converted to text +
    structured output validated by Pydantic."""
    import anthropic

    # max_retries=0: one real call takes ~46 s (measured on case 01); a retry would never fit
    # the 60 s serverless ceiling. Better to fail fast and let the lawyer resend.
    client = anthropic.Anthropic(timeout=TIMEOUT_SECONDS, max_retries=0)
    model = os.environ.get("EXTRACTION_MODEL", DEFAULT_MODEL)

    content: list[dict] = []
    for name, data in files:
        if _is_pdf(data):
            content.append(
                {
                    "type": "document",
                    "source": {
                        "type": "base64",
                        "media_type": "application/pdf",
                        "data": base64.standard_b64encode(data).decode(),
                    },
                    "title": name,
                }
            )
        else:
            content.append({"type": "text", "text": spreadsheet_to_text(name, data)})
    content.append(
        {
            "type": "text",
            "text": "Extract the calculation parameters from these court documents (they may be of different "
            "kinds: judgment, appellate decision, motion for clarification, settlement approval, payment "
            "history, contract, spreadsheet...). Classify each document, apply the decision hierarchy and "
            "follow the system rules strictly.",
        }
    )

    response = client.messages.parse(
        model=model,
        max_tokens=8000,
        # thinking DISABLED on purpose: the extraction is a literal transport of the text, not
        # deliberation. Measured on case 01 (2 PDFs): with adaptive thinking on, the model spent
        # the WHOLE 8000 tokens thinking (91 s, stop max_tokens, truncated JSON); disabled:
        # 46 s, stop end_turn, 7/7 fields of the answer key correct.
        thinking={"type": "disabled"},
        system=[{"type": "text", "text": INSTRUCTIONS, "cache_control": {"type": "ephemeral"}}],
        messages=[{"role": "user", "content": content}],
        output_format=ExtractionResult,
    )
    extracted = _normalise(response.parsed_output)
    return {
        "simulated": False,
        "model": model,
        "extraction": extracted.model_dump(),
        "usage": {
            "input_tokens": response.usage.input_tokens,
            "output_tokens": response.usage.output_tokens,
        },
    }


def extract(files: list[tuple[str, bytes]]) -> dict:
    if extraction_mode() == "live":
        return extract_with_model(files)
    return simulated_extraction([name for name, _ in files])
